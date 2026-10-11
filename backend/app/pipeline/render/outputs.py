"""Word (.docx) and plain-text outputs (PLAN.md §6.1, stage 7)."""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

from ...models import Job, Page, Segment
from ..normalize import strip_control

RED = RGBColor(0xB4, 0x23, 0x18)
GREY = RGBColor(0x6B, 0x72, 0x80)


def _paragraphs(segments: list[Segment]) -> dict[int, dict[int, list[Segment]]]:
    """page_no -> para_no -> segments in order."""
    pages: dict[int, dict[int, list[Segment]]] = defaultdict(lambda: defaultdict(list))
    for s in sorted(segments, key=lambda s: s.order):
        pages[s.page_no][s.para_no].append(s)
    return pages


def _is_untranslated(s: Segment) -> bool:
    return s.edited is None and s.chosen_engine == "untranslated"


def output_name(job: Job, ext: str) -> str:
    stem = Path(job.filename).stem or "document"
    return f"{stem}_English.{ext}"


def render_docx(job: Job, pages: list[Page], segments: list[Segment], path: Path, model_label: str = "") -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    props = doc.core_properties
    props.title = f"{Path(job.filename).stem} (English)"
    props.comments = (
        f"Translated by: {model_label or job.model_used or job.draft_model or 'n/a'}; "
        f"quick draft: {job.draft_model or 'n/a'}; generated {datetime.now():%Y-%m-%d %H:%M}"
    )

    doc.add_heading(Path(job.filename).stem, level=1)
    meta = doc.add_paragraph()
    run = meta.add_run(f"English translation · {model_label or job.model_used or job.draft_model or ''}")
    run.font.size = Pt(9)
    run.font.color.rgb = GREY

    by_page = _paragraphs(segments)
    page_notes = {p.page_no: p.note for p in pages}
    page_numbers = sorted({p.page_no for p in pages} | set(by_page))
    for i, page_no in enumerate(page_numbers):
        if len(page_numbers) > 1:
            h = doc.add_paragraph()
            r = h.add_run(f"Page {page_no}")
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = GREY
        paras = by_page.get(page_no, {})
        if not paras and page_notes.get(page_no):
            r = doc.add_paragraph().add_run(page_notes[page_no])
            r.italic = True
            r.font.color.rgb = GREY
        for para_no in sorted(paras):
            p = doc.add_paragraph()
            for j, seg in enumerate(paras[para_no]):
                text = strip_control(seg.display_text or seg.source)
                if j:
                    p.add_run(" ")
                if _is_untranslated(seg):
                    r = p.add_run(f"[needs review: {text}]")
                    r.font.color.rgb = RED
                else:
                    p.add_run(text)
        if i < len(page_numbers) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def render_txt(job: Job, pages: list[Page], segments: list[Segment], path: Path) -> Path:
    lines: list[str] = []
    by_page = _paragraphs(segments)
    page_numbers = sorted({p.page_no for p in pages} | set(by_page))
    for page_no in page_numbers:
        if len(page_numbers) > 1:
            lines.append(f"--- Page {page_no} ---")
        for para_no in sorted(by_page.get(page_no, {})):
            parts = []
            for seg in by_page[page_no][para_no]:
                text = seg.display_text or seg.source
                parts.append(f"[needs review: {text}]" if _is_untranslated(seg) else text)
            lines.append(" ".join(parts))
            lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path
