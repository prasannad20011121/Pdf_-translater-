"""End-to-end: upload -> worker stages -> review edits -> downloads, using the dev stand-in."""

import io
import shutil

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.pipeline.ocr import tesseract

HAS_OCR = tesseract.available(shutil.which("tesseract"))


@pytest.fixture()
def client(app_env):
    from app.main import app

    with TestClient(app) as c:
        yield c


def upload(client, path, choice="auto"):
    with open(path, "rb") as f:
        r = client.post("/api/jobs", files={"file": (path.name, f, "application/pdf")}, data={"choice": choice})
    assert r.status_code == 200, r.text
    return r.json()


def run_worker(app_env, job_id):
    from app.worker import JobRunner

    JobRunner(app_env).run(job_id)


def test_settings_report_dev_profile(client):
    data = client.get("/api/settings").json()
    assert data["dev_mode"] is True
    assert {c["id"] for c in data["choices"]} == {"auto", "fast", "balanced", "best"}


def test_clean_pdf_goes_through_all_stages(client, app_env, pdfs):
    job = upload(client, pdfs["clean"])
    assert job["status"] == "queued" and job["page_count"] == 1
    run_worker(app_env, job["id"])

    job = client.get(f"/api/jobs/{job['id']}").json()
    assert job["status"] == "done", job
    assert job["draft_ready"] and job["final_ready"]
    assert job["draft_model"] == "dev-standin-draft"
    assert job["model_used"] == "dev-standin-llm"
    assert job["pages"][0]["kind"] == "unicode" and job["pages"][0]["method"] == "native"

    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert len(segs) >= 6
    memo = next(s for s in segs if "ज्ञापांक" in s["source"])
    assert "1234/2026" in memo["text"] and "12.03.2026" in memo["text"]
    assert memo["engine"] == "dev-standin-llm"
    assert all(s["confidence"] in ("green", "amber") for s in segs)
    assert sum(job["counts"].values()) == len(segs)

    r = client.get(f"/api/jobs/{job['id']}/download/docx")
    assert r.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "1234/2026" in text
    assert client.get(f"/api/jobs/{job['id']}/download/txt").status_code == 200
    assert client.get(f"/api/jobs/{job['id']}/pages/1.png").headers["content-type"] == "image/png"


def test_fast_choice_uses_draft_only(client, app_env, pdfs):
    job = upload(client, pdfs["clean"], choice="fast")
    run_worker(app_env, job["id"])
    job = client.get(f"/api/jobs/{job['id']}").json()
    assert job["status"] == "done" and job["model_used"] is None
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert {s["engine"] for s in segs if s["needs_translation"]} == {"dev-standin-draft"}


def test_english_page_is_copied(client, app_env, pdfs):
    job = upload(client, pdfs["english"])
    run_worker(app_env, job["id"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert segs and all(s["engine"] == "copy" and s["text"] == s["source"] for s in segs)


@pytest.mark.skipif(not HAS_OCR, reason="tesseract with Hindi not installed")
def test_scanned_pdf_uses_ocr_and_is_flagged(client, app_env, pdfs):
    job = upload(client, pdfs["scan"])
    run_worker(app_env, job["id"])
    job = client.get(f"/api/jobs/{job['id']}").json()
    page = job["pages"][0]
    assert page["kind"] == "scan" and page["method"] == "ocr-tesseract" and page["flagged"]
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert segs and all(s["confidence"] != "green" for s in segs if s["needs_translation"])


@pytest.mark.skipif(not HAS_OCR, reason="tesseract with Hindi not installed")
def test_garbled_text_layer_is_replaced_by_ocr(client, app_env, pdfs):
    job = upload(client, pdfs["broken"])
    run_worker(app_env, job["id"])
    job = client.get(f"/api/jobs/{job['id']}").json()
    assert job["pages"][0]["method"] == "ocr-tesseract"
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert any("निविदा" in s["source"] for s in segs)


def test_correction_feeds_translation_memory(client, app_env, pdfs):
    job = upload(client, pdfs["clean"])
    run_worker(app_env, job["id"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    target = next(s for s in segs if s["source"].startswith("आदेश से"))

    before = client.get(f"/api/jobs/{job['id']}").json()["counts"]
    r = client.patch(f"/api/segments/{target['id']}", json={"text": "By order, Managing Director"})
    assert r.json()["text"] == "By order, Managing Director" and r.json()["edited"]
    after = client.get(f"/api/jobs/{job['id']}").json()["counts"]
    assert sum(after.values()) == sum(before.values())
    assert after["green"] >= before["green"]
    assert client.patch(f"/api/segments/{target['id']}", json={"text": "   "}).status_code == 400

    docx = client.get(f"/api/jobs/{job['id']}/download/docx").content
    assert "By order, Managing Director" in "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)

    # A second document reuses the corrected sentence from translation memory.
    job2 = upload(client, pdfs["clean"])
    run_worker(app_env, job2["id"])
    segs2 = client.get(f"/api/jobs/{job2['id']}/segments").json()
    again = next(s for s in segs2 if s["source"].startswith("आदेश से"))
    assert again["engine"] == "memory" and again["text"] == "By order, Managing Director"


def test_glossary_crud_and_csv(client):
    r = client.post("/api/glossary", json={"hindi": "निविदा", "english": "Tender"})
    assert r.status_code == 200
    csv_data = "hindi,english,note\nकार्यादेश,Work Order,\nज्ञापांक,Memo No.,\n".encode()
    r = client.post("/api/glossary/import", files={"file": ("g.csv", csv_data, "text/csv")})
    assert r.json() == {"imported": 2, "skipped": 0, "errors": []}
    terms = client.get("/api/glossary").json()
    assert {t["hindi"] for t in terms} == {"निविदा", "कार्यादेश", "ज्ञापांक"}
    exported = client.get("/api/glossary/export").content.decode("utf-8-sig")
    assert "Work Order" in exported
    client.delete(f"/api/glossary/{terms[0]['id']}")
    assert len(client.get("/api/glossary").json()) == 2
    assert client.get("/api/tips").json()["term"] is not None


def test_upload_rejects_non_pdf_and_encrypted(client, tmp_path, pdfs):
    r = client.post("/api/jobs", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert r.status_code == 400
    r = client.post("/api/jobs", files={"file": ("fake.pdf", b"not a pdf", "application/pdf")})
    assert r.status_code == 400

    import pymupdf

    doc = pymupdf.open(pdfs["clean"])
    locked = tmp_path / "locked.pdf"
    doc.save(str(locked), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    with open(locked, "rb") as f:
        r = client.post("/api/jobs", files={"file": ("locked.pdf", f, "application/pdf")})
    assert r.status_code == 400 and "Password" in r.json()["detail"]


def test_resume_after_interruption(client, app_env, pdfs):
    """A job interrupted after extraction carries on without duplicating sentences."""
    from app.db import session_scope
    from app.models import Job
    from app.worker import JobRunner

    job = upload(client, pdfs["clean"])
    runner = JobRunner(app_env)
    runner.stage_extract(job["id"])
    n = len(client.get(f"/api/jobs/{job['id']}/segments").json())
    with session_scope() as db:
        j = db.get(Job, job["id"])
        j.status, j.stage = "running", "draft"
    runner.run(job["id"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert len(segs) == n
    assert client.get(f"/api/jobs/{job['id']}").json()["status"] == "done"


def test_models_endpoint_lists_catalog(client):
    models = client.get("/api/models").json()
    ids = {m["id"] for m in models}
    assert {"gemma4-e2b", "gemma4-e4b", "indictrans2-indic-en-1b"} <= ids
    gemma = next(m for m in models if m["id"] == "gemma4-e2b")
    assert gemma["installed"] is False and gemma["licence"] == "Apache-2.0"
