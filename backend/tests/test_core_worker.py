"""Worker robustness: delete/cancel, engine failures, broken pages, memory, retry,
counts, rendering, the single-worker lock and graceful stop."""

from __future__ import annotations

import io
import json
import os
import signal
import stat
import subprocess
import sys
import threading
import time
from datetime import timedelta
from pathlib import Path

import psutil
import pymupdf
import pytest
from app.model_manager import ModelManager
from app.pipeline.translate.base import Translator
from conftest import LETTER, make_clean_pdf
from docx import Document
from fastapi.testclient import TestClient

BACKEND = Path(__file__).resolve().parents[1]
DRAFT, LLM = "dev-standin-draft", "dev-standin-llm"


class Hooked(Translator):
    """Wraps a real engine; records calls and runs test hooks (which may raise)."""

    def __init__(self, inner: Translator, events: list, hooks: dict):
        self.inner, self.events, self.hooks = inner, events, hooks
        self.id, self.label = inner.id, inner.label
        self.unit, self.uses_placeholders = inner.unit, inner.uses_placeholders

    def _hook(self, action: str) -> None:
        self.events.append((action, self.id))
        hook = self.hooks.get((action, self.id))
        if hook is not None:
            hook()

    def load(self) -> None:
        self._hook("load")
        self.inner.load()

    def unload(self) -> None:
        self.events.append(("unload", self.id))
        self.inner.unload()

    def translate(self, texts, ctx):
        self._hook("translate")
        return self.inner.translate(texts, ctx)


class HookedManager(ModelManager):
    def __init__(self, settings):
        super().__init__(settings)
        self.events: list[tuple[str, str]] = []
        self.hooks: dict = {}

    def create(self, model_id):
        return Hooked(super().create(model_id), self.events, self.hooks)

    def balanced(self) -> bool:
        """Every engine that was loaded was also unloaded."""
        for model_id in {m for _, m in self.events}:
            loads = sum(1 for a, m in self.events if a == "load" and m == model_id)
            unloads = sum(1 for a, m in self.events if a == "unload" and m == model_id)
            if unloads < loads:
                return False
        return True


@pytest.fixture()
def client(app_env):
    from app.main import app

    with TestClient(app) as c:
        yield c


@pytest.fixture()
def manager(app_env):
    return HookedManager(app_env)


@pytest.fixture()
def runner(app_env, manager):
    from app.worker import JobRunner

    return JobRunner(app_env, manager)


def upload(client, path, choice="auto"):
    with open(path, "rb") as f:
        r = client.post("/api/jobs", files={"file": (path.name, f, "application/pdf")}, data={"choice": choice})
    assert r.status_code == 200, r.text
    return r.json()


def job_row(job_id):
    from app.db import session_scope
    from app.models import Job

    with session_scope() as db:
        return db.get(Job, job_id)


def two_page_pdf(path: Path) -> Path:
    one = make_clean_pdf(path.with_suffix(".one.pdf"), LETTER)
    doc = pymupdf.open(str(one))
    doc.insert_pdf(pymupdf.open(str(one)))
    doc.save(str(path))
    return path


# ---- delete / cancel ------------------------------------------------------------------

def test_delete_running_job_cancels_it_and_the_worker_cleans_up(client, app_env, pdfs, manager, runner):
    job = upload(client, pdfs["clean"])
    seen = {}

    def delete_now():
        r = client.delete(f"/api/jobs/{job['id']}")
        seen["delete"] = r.json()
        seen["listed"] = [j["id"] for j in client.get("/api/jobs").json()]
        seen["get"] = client.get(f"/api/jobs/{job['id']}").status_code

    manager.hooks[("translate", DRAFT)] = delete_now
    runner.run(job["id"])  # must not raise

    assert seen["delete"] == {"status": "cancelling"}
    assert job["id"] not in seen["listed"] and seen["get"] == 404
    assert job_row(job["id"]) is None
    assert not app_env.job_dir(job["id"]).exists()
    assert manager.balanced()
    assert ("load", LLM) not in manager.events  # stopped before the final stage

    # The same worker carries on with the next document.
    manager.hooks.clear()
    job2 = upload(client, pdfs["clean"])
    runner.run(job2["id"])
    assert client.get(f"/api/jobs/{job2['id']}").json()["status"] == "done"


def test_delete_while_final_engine_starts_still_unloads_it(client, app_env, pdfs, manager, runner):
    job = upload(client, pdfs["clean"])
    manager.hooks[("load", LLM)] = lambda: client.delete(f"/api/jobs/{job['id']}")
    runner.run(job["id"])
    assert ("load", LLM) in manager.events and ("unload", LLM) in manager.events
    assert manager.balanced()
    assert job_row(job["id"]) is None


def test_row_deleted_under_the_worker_does_not_crash_it(client, app_env, pdfs, manager, runner):
    from app.worker import remove_job

    job = upload(client, pdfs["clean"])
    manager.hooks[("translate", DRAFT)] = lambda: remove_job(app_env, job["id"])
    runner.run(job["id"])  # StaleDataError / missing row: return quietly
    assert job_row(job["id"]) is None and manager.balanced()


def test_delete_queued_job_is_hidden_and_purged_by_the_worker(client, app_env, pdfs):
    from app.worker import next_job_id, purge_cancelled

    job = upload(client, pdfs["clean"])
    assert client.delete(f"/api/jobs/{job['id']}").json() == {"status": "cancelling"}
    assert client.delete(f"/api/jobs/{job['id']}").json() == {"status": "cancelling"}  # idempotent
    assert client.get("/api/jobs").json() == []
    assert client.get(f"/api/jobs/{job['id']}").status_code == 404
    assert next_job_id() is None
    assert purge_cancelled(app_env) == 1
    assert job_row(job["id"]) is None and not app_env.job_dir(job["id"]).exists()


def test_delete_finished_job_is_immediate(client, app_env, pdfs, runner):
    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    assert client.delete(f"/api/jobs/{job['id']}").json() == {"status": "deleted"}
    assert job_row(job["id"]) is None and not app_env.job_dir(job["id"]).exists()
    assert client.get(f"/api/jobs/{job['id']}").status_code == 404
    assert client.delete(f"/api/jobs/{job['id']}").status_code == 404


def test_cleanup_never_touches_queued_or_running_jobs(client, app_env, pdfs):
    from app.db import session_scope
    from app.models import Job, utcnow
    from app.worker import cleanup_old_jobs

    old_running, old_done, new_done = (upload(client, pdfs["clean"])["id"] for _ in range(3))
    with session_scope() as db:
        for job_id, status, age in ((old_running, "running", 40), (old_done, "done", 40), (new_done, "done", 1)):
            j = db.get(Job, job_id)
            j.status, j.created_at = status, utcnow() - timedelta(days=age)
    assert cleanup_old_jobs(app_env) == 1
    assert job_row(old_done) is None and not app_env.job_dir(old_done).exists()
    assert job_row(old_running) is not None and (app_env.job_dir(old_running) / "source.pdf").exists()
    assert job_row(new_done) is not None


def test_cleanup_removes_old_orphan_folders_only(app_env):
    from app.worker import cleanup_old_jobs

    old, fresh = app_env.job_dir("a" * 32), app_env.job_dir("b" * 32)
    for d in (old, fresh):
        d.mkdir(parents=True)
    past = time.time() - 7200
    os.utime(old, (past, past))
    cleanup_old_jobs(app_env)
    assert not old.exists() and fresh.exists()  # a fresh folder may be an upload in progress


# ---- engine failures are fallbacks, never failed jobs ------------------------------------

def test_draft_engine_load_error_falls_back_to_final(client, pdfs, manager, runner):
    job = upload(client, pdfs["clean"])

    def blocked():
        raise OSError("[WinError 1260] This program is blocked by group policy")

    manager.hooks[("load", DRAFT)] = blocked
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done" and data["model_used"] == LLM
    assert any("could not start" in n for n in data["notes"])
    assert manager.balanced()


def test_final_engine_load_error_keeps_the_draft_amber(client, pdfs, manager, runner):
    job = upload(client, pdfs["clean"])

    def denied():
        raise PermissionError(13, "Access is denied")

    manager.hooks[("load", LLM)] = denied
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done", data
    assert data["model_used"] is None and any("could not start" in n for n in data["notes"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    drafted = [s for s in segs if s["engine"] == DRAFT]
    assert drafted and all(s["confidence"] != "green" for s in drafted)
    assert all("fallback" in {i["code"] for i in s["issues"]} for s in drafted)
    assert manager.balanced()


def test_translate_errors_fall_back(client, pdfs, manager, runner):
    job = upload(client, pdfs["clean"])

    def oom():
        raise RuntimeError("DefaultCPUAllocator: not enough memory")

    manager.hooks[("translate", DRAFT)] = oom
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done" and data["model_used"] == LLM
    assert any("stopped" in n for n in data["notes"])

    job = upload(client, pdfs["clean"])
    manager.hooks.clear()
    manager.hooks[("translate", LLM)] = oom
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done"
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert {s["engine"] for s in segs if s["needs_translation"]} == {DRAFT}
    assert manager.balanced()


def test_engine_returning_too_few_outputs_does_not_loop(client, pdfs, manager, runner, monkeypatch):
    from app.pipeline.translate.devstub import DevStandIn

    real = DevStandIn.translate
    monkeypatch.setattr(DevStandIn, "translate", lambda self, texts, ctx: real(self, texts, ctx)[:1])
    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    assert client.get(f"/api/jobs/{job['id']}").json()["status"] == "done"


# ---- extraction ----------------------------------------------------------------------------

def test_broken_page_is_skipped_with_a_note(client, app_env, tmp_path, manager, runner, monkeypatch):
    from app import worker

    real = worker.classify_page

    def flaky(page):
        if page.number == 1:
            raise RuntimeError("MuPDF error: invalid content stream")
        return real(page)

    monkeypatch.setattr(worker, "classify_page", flaky)
    job = upload(client, two_page_pdf(tmp_path / "two.pdf"))
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done", data
    pages = {p["page_no"]: p for p in data["pages"]}
    assert pages[2]["flagged"] and "could not be read" in pages[2]["note"]
    assert pages[1]["method"] == "native"
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert segs and {s["page_no"] for s in segs} == {1}
    assert any("Page 2" in n for n in data["notes"])


def test_resume_after_a_broken_page_keeps_sentence_order_unique(client, app_env, tmp_path, runner, monkeypatch):
    """Sentence numbers continue after the highest one, even with gaps."""
    from app import worker
    from app.db import session_scope
    from app.models import Job, Segment

    job = upload(client, two_page_pdf(tmp_path / "two.pdf"))
    real = worker.JobRunner._extract_page
    calls = {"n": 0}

    def stop_after_first(self, job_id, doc, index, *args):
        calls["n"] += 1
        if index == 1:
            raise KeyboardInterrupt  # the laptop "shuts down" after page 1
        return real(self, job_id, doc, index, *args)

    monkeypatch.setattr(worker.JobRunner, "_extract_page", stop_after_first)
    with pytest.raises(KeyboardInterrupt):
        runner.run(job["id"])
    monkeypatch.setattr(worker.JobRunner, "_extract_page", real)
    with session_scope() as db:
        db.get(Job, job["id"]).status = "running"
    runner.run(job["id"])
    with session_scope() as db:
        orders = [s.order for s in db.query(Segment).filter(Segment.job_id == job["id"])]
    assert len(orders) == len(set(orders)) and client.get(f"/api/jobs/{job['id']}").json()["status"] == "done"


def test_unreadable_pdf_fails_clearly_and_can_be_retried(client, app_env, pdfs, runner):
    job = upload(client, pdfs["clean"])
    source = app_env.job_dir(job["id"]) / "source.pdf"
    good = source.read_bytes()
    source.write_bytes(b"%PDF-1.7 truncated")
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "failed" and "could not be opened" in data["error"]

    source.write_bytes(good)
    r = client.post(f"/api/jobs/{job['id']}/retry")
    assert r.status_code == 200 and r.json()["status"] == "queued" and r.json()["error"] is None
    assert client.post(f"/api/jobs/{job['id']}/retry").status_code == 409  # no longer failed
    runner.run(job["id"])
    assert client.get(f"/api/jobs/{job['id']}").json()["status"] == "done"


def test_retry_errors(client):
    assert client.post(f"/api/jobs/{'0' * 32}/retry").status_code == 404
    assert client.post("/api/jobs/not-an-id/retry").status_code == 404


def _legacy_pdf(path: Path, letterhead_only: bool) -> Path:
    """Text in a font named KrutiDev010: it copies out as Latin gibberish."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((60, 72), "fcgkj ljdkj lwpuk izkS|ksfxdh foHkkx", fontname="helv", fontsize=12)
    raw = pymupdf.open(stream=doc.tobytes(), filetype="pdf")
    for xref in range(1, raw.xref_length()):
        if raw.xref_get_key(xref, "Type")[1] == "/Font":
            raw.xref_set_key(xref, "BaseFont", "/KrutiDev010")
    if letterhead_only:
        # The rest of the letter is normal Unicode Hindi (well over 70% of the text).
        body = make_clean_pdf(path.with_suffix(".body.pdf"), LETTER)
        page = raw[0]
        page.show_pdf_page(pymupdf.Rect(0, 100, page.rect.width, page.rect.height), pymupdf.open(str(body)), 0)
    raw.save(str(path))
    return path


def test_legacy_font_page_without_ocr_is_never_copied_as_english(client, app_env, tmp_path, runner, monkeypatch):
    monkeypatch.setattr(app_env, "tesseract_cmd", None)
    job = upload(client, _legacy_pdf(tmp_path / "kruti.pdf", letterhead_only=False))
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    page = data["pages"][0]
    assert page["kind"] == "legacy" and page["flagged"] and "old Hindi font" in page["note"]
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert not any("fcgkj" in (s["text"] or "") and s["confidence"] == "green" for s in segs)
    assert segs == []


def test_legacy_font_letterhead_is_never_copied_as_english(client, app_env, tmp_path, runner, monkeypatch):
    """A Kruti Dev letterhead on a Unicode letter: the gibberish is either left out or
    red/untranslated (depending on how the page is classified), never copied as green."""
    monkeypatch.setattr(app_env, "tesseract_cmd", None)
    job = upload(client, _legacy_pdf(tmp_path / "letterhead.pdf", letterhead_only=True))
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    gibberish = [s for s in segs if "fcgkj" in s["source"]]
    assert all(s["confidence"] == "red" and s["engine"] == "untranslated" for s in gibberish)
    assert any("ज्ञापांक" in s["source"] for s in segs)  # the Unicode body is still translated
    assert data["counts"]["red"] >= len(gibberish)
    assert data["pages"][0]["note"]  # the reader is told


def test_legacy_page_without_ocr_keeps_only_unicode_hindi(client, app_env, tmp_path, runner, monkeypatch):
    """A page marked legacy because of a Kruti Dev letterhead still gives its Unicode body."""
    from app import worker
    from app.pipeline.classify import PageInfo

    monkeypatch.setattr(app_env, "tesseract_cmd", None)
    monkeypatch.setattr(worker, "classify_page", lambda page: PageInfo(kind="legacy", chars=500))
    job = upload(client, _legacy_pdf(tmp_path / "letterhead.pdf", letterhead_only=True))
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    page = data["pages"][0]
    assert page["flagged"] and "left out" in page["note"]
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert segs and not any("fcgkj" in s["source"] for s in segs)
    assert any("ज्ञापांक" in s["source"] for s in segs)


def test_numbers_the_ocr_was_unsure_of_make_the_sentence_red(client, app_env, pdfs, runner, monkeypatch):
    from types import SimpleNamespace

    from app import worker
    from app.pipeline.classify import PageInfo

    blocks = [
        SimpleNamespace(text="ज्ञापांक-1234/2026, दिनांक 12.03.2026", bbox=None, uncertain_numbers=["1234/2026"]),
        SimpleNamespace(text="Rs. 11,800/-", bbox=None, uncertain_numbers=["11,800"]),
        SimpleNamespace(text="निविदा की अंतिम तिथि बढ़ाई गई है।", bbox=None, uncertain_numbers=[]),
    ]
    monkeypatch.setattr(app_env, "tesseract_cmd", "tesseract")
    monkeypatch.setattr(worker.tesseract, "available", lambda cmd: True)
    monkeypatch.setattr(worker.tesseract, "ocr_image",
                        lambda image, cmd, **kw: SimpleNamespace(blocks=blocks, mean_confidence=90.0))
    monkeypatch.setattr(worker, "classify_page", lambda page: PageInfo(kind="scan", chars=0))
    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    by_source = {s["source"]: s for s in segs}
    memo = next(s for src, s in by_source.items() if "1234/2026" in src)
    amount = by_source["Rs. 11,800/-"]
    plain = next(s for src, s in by_source.items() if src.startswith("निविदा"))
    for s in (memo, amount):
        assert s["confidence"] == "red" and "ocr_numbers" in {i["code"] for i in s["issues"]}, s
    assert memo["engine"] == LLM  # still translated, just marked for checking
    assert plain["confidence"] == "amber" and "ocr_numbers" not in {i["code"] for i in plain["issues"]}


def test_copied_text_on_flagged_page_is_amber(client, app_env, pdfs, runner, monkeypatch):
    """Numbers and English read by OCR (or from a garbled text layer) are never green."""
    from app import worker
    from app.pipeline.classify import PageInfo

    monkeypatch.setattr(app_env, "tesseract_cmd", None)
    # A garbled Unicode page with no OCR: native text is used but the page is flagged.
    monkeypatch.setattr(worker, "classify_page",
                        lambda page: PageInfo(kind="unicode", chars=100, quality=0.1))
    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    copied = [s for s in segs if s["engine"] == "copy"]
    assert segs and all(s["confidence"] != "green" for s in segs)
    for s in copied:
        assert s["confidence"] == "amber" and "ocr" in {i["code"] for i in s["issues"]}


# ---- memory and "nothing could translate" ---------------------------------------------------

def _no_memory(monkeypatch, draft_too: bool = True):
    from app import model_manager, worker

    monkeypatch.setattr(worker, "fits_in_memory", lambda gb: False)
    if draft_too:
        monkeypatch.setattr(model_manager, "fits_in_memory", lambda gb: False)


def test_no_free_memory_fails_with_an_actionable_error_and_retry_works(client, pdfs, manager, runner, monkeypatch):
    job = upload(client, pdfs["clean"])
    _no_memory(monkeypatch)
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "failed" and data["stage"] == "draft"
    assert "Close other apps" in data["error"] and "Try again" in data["error"]
    assert not data["final_ready"] and not data["draft_ready"]
    assert not any("used the quick draft" in n for n in data["notes"])
    assert manager.events == []  # nothing was loaded

    monkeypatch.undo()
    r = client.post(f"/api/jobs/{job['id']}/retry")
    assert r.status_code == 200 and r.json()["notes"] == []
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done" and data["model_used"] == LLM


def test_memory_recovers_while_waiting(client, pdfs, runner, monkeypatch):
    from app import model_manager

    state = {"calls": 0}

    def fits(gb):
        state["calls"] += 1
        return state["calls"] > 2  # the user closed Chrome

    monkeypatch.setattr(model_manager, "fits_in_memory", fits)
    job = upload(client, pdfs["clean"], choice="fast")
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done" and data["draft_model"] == DRAFT


def test_nothing_installed_fails_with_install_hint(client, pdfs, runner, monkeypatch):
    from app.model_manager import Readiness

    monkeypatch.setattr(runner.manager, "readiness", lambda model_id: Readiness(False, "not installed"))
    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "failed" and "Settings > Models" in data["error"]


def test_final_model_skipped_for_memory_marks_drafts_amber(client, pdfs, runner, monkeypatch):
    _no_memory(monkeypatch, draft_too=False)  # IndicTrans2 fits, Gemma doesn't
    job = upload(client, pdfs["clean"], choice="balanced")
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done" and data["model_used"] is None
    assert any("not enough free memory" in n for n in data["notes"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    drafted = [s for s in segs if s["engine"] == DRAFT]
    assert drafted and all(s["confidence"] in ("amber", "red") for s in drafted)
    assert data["counts"]["green"] < len(segs)


def test_fast_choice_drafts_can_be_green(client, pdfs, runner):
    job = upload(client, pdfs["clean"], choice="fast")
    runner.run(job["id"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    assert any(s["confidence"] == "green" and s["engine"] == DRAFT for s in segs)


def test_no_model_is_loaded_when_nothing_needs_translating(client, pdfs, manager, runner):
    job = upload(client, pdfs["english"])
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done" and data["draft_ready"]
    assert data["model_used"] is None and data["draft_model"] is None
    assert manager.events == []


def test_eta_only_counts_a_final_stage_that_will_run(client, app_env, pdfs, monkeypatch):
    from app import worker

    def eta_after_draft(choice):
        job = upload(client, pdfs["clean"], choice=choice)
        r = worker.JobRunner(app_env)
        r.stage_extract(job["id"])
        r.stage_draft(job["id"])
        return job_row(job["id"]).eta_seconds

    assert eta_after_draft("auto") > 0  # the final model will run after the draft
    monkeypatch.setattr(worker, "fits_in_memory", lambda gb: False)
    assert eta_after_draft("auto") == 0  # it won't fit: no phantom final-stage time
    assert eta_after_draft("fast") == 0


def test_speed_is_saved_per_model(client, app_env, pdfs, runner):
    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    rates = json.loads((app_env.home / "speed.json").read_text())
    assert f"draft:{DRAFT}" in rates and f"final:{LLM}" in rates


# ---- counts ----------------------------------------------------------------------------------

def test_counts_stay_right_during_and_after_the_run(client, app_env, pdfs, runner):
    job = upload(client, pdfs["clean"])
    runner.stage_extract(job["id"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    counts = client.get(f"/api/jobs/{job['id']}").json()["counts"]
    # Untranslated sentences are not counted as good.
    assert sum(counts.values()) == sum(1 for s in segs if s["confidence"])
    assert sum(counts.values()) < len(segs)

    runner.stage_draft(job["id"])
    target = next(s for s in client.get(f"/api/jobs/{job['id']}/segments").json() if s["needs_translation"])
    client.patch(f"/api/segments/{target['id']}", json={"text": "Corrected by a person."})
    # Make the rest of the machine output red, so a lost correction would show.
    from app.db import session_scope
    from app.models import Job

    with session_scope() as db:
        db.get(Job, job["id"]).stage = "final"
        db.get(Job, job["id"]).status = "running"
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    expected = {"green": 0, "amber": 0, "red": 0}
    for s in segs:
        expected["green" if s["edited"] else s["confidence"]] += 1
    assert data["status"] == "done" and data["counts"] == expected
    assert next(s for s in segs if s["id"] == target["id"])["edited"]


# ---- rendering -------------------------------------------------------------------------------

def test_a_correction_made_during_a_render_is_not_lost(client, app_env, pdfs, runner, monkeypatch):
    from app import worker

    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    segs = client.get(f"/api/jobs/{job['id']}/segments").json()
    first, second = [s for s in segs if s["needs_translation"]][:2]
    client.patch(f"/api/segments/{first['id']}", json={"text": "FIRST CORRECTION"})

    real = worker.render_docx
    done = {}

    def render_and_edit(*args, **kwargs):
        path = real(*args, **kwargs)
        if not done:
            done["x"] = client.patch(f"/api/segments/{second['id']}", json={"text": "SECOND CORRECTION"}).status_code
        return path

    monkeypatch.setattr(worker, "render_docx", render_and_edit)
    client.get(f"/api/jobs/{job['id']}/download/docx")
    assert done["x"] == 200
    assert job_row(job["id"]).outputs_stale  # the render saw the old text
    docx = client.get(f"/api/jobs/{job['id']}/download/docx").content
    text = "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)
    assert "FIRST CORRECTION" in text and "SECOND CORRECTION" in text


def test_text_file_is_written_even_if_word_fails(client, app_env, pdfs, runner, monkeypatch):
    from app import worker

    def broken(*args, **kwargs):
        raise ValueError("All strings must be XML compatible")

    monkeypatch.setattr(worker, "render_docx", broken)
    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    data = client.get(f"/api/jobs/{job['id']}").json()
    assert data["status"] == "done", data
    assert any("Word file could not be created" in n for n in data["notes"])
    assert client.get(f"/api/jobs/{job['id']}/download/txt").status_code == 200
    r = client.get(f"/api/jobs/{job['id']}/download/docx")
    assert r.status_code == 500 and "text file" in r.json()["detail"]


def test_characters_word_rejects_are_cleaned(client, app_env, pdfs, runner):
    job = upload(client, pdfs["clean"])
    runner.run(job["id"])
    seg = next(s for s in client.get(f"/api/jobs/{job['id']}/segments").json() if s["needs_translation"])
    client.patch(f"/api/segments/{seg['id']}", json={"text": "Broken\uffff text\ufffe here"})
    r = client.get(f"/api/jobs/{job['id']}/download/docx")
    assert r.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "Broken text here" in text


def test_outputs_have_short_fixed_names(client, app_env, pdfs, runner):
    long_name = pdfs["clean"].with_name("x" * 200 + ".pdf")
    long_name.write_bytes(pdfs["clean"].read_bytes())
    job = upload(client, long_name)
    runner.run(job["id"])
    out = app_env.job_dir(job["id"]) / "out"
    assert sorted(p.name for p in out.iterdir()) == ["translation.docx", "translation.txt"]
    r = client.get(f"/api/jobs/{job['id']}/download/docx")
    assert r.status_code == 200 and ("x" * 50 + "") in r.headers["content-disposition"]


# ---- the worker loop, the lock and stopping ---------------------------------------------------

def test_run_forever_survives_a_job_that_crashes(client, app_env, pdfs, monkeypatch):
    from app import worker

    bad, good = upload(client, pdfs["clean"]), upload(client, pdfs["clean"])
    real = worker.JobRunner.run

    def crashing(self, job_id):
        if job_id == bad["id"]:
            raise AttributeError("'NoneType' object has no attribute 'status'")
        return real(self, job_id)

    monkeypatch.setattr(worker.JobRunner, "run", crashing)
    monkeypatch.setattr(worker, "_wait", lambda stop, seconds: stop.wait(0.01))
    stop = threading.Event()
    t = threading.Thread(target=worker.run_forever, kwargs={"poll_seconds": 0.01, "stop": stop})
    t.start()
    try:
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and job_row(good["id"]).status != "done":
            time.sleep(0.05)
    finally:
        stop.set()
        t.join(10)
    assert job_row(good["id"]).status == "done"
    assert job_row(bad["id"]).status == "failed"  # given up after repeated crashes, not looping forever


def test_only_one_worker_can_run(app_env):
    from app.worker import LOCK_EXIT_CODE, WorkerLock, run_forever

    first = WorkerLock(app_env.home)
    assert first.acquire()
    try:
        assert first.holder() == os.getpid()
        assert not WorkerLock(app_env.home).acquire()
        assert run_forever(poll_seconds=0.01, stop=threading.Event()) == LOCK_EXIT_CODE
    finally:
        first.release()
    again = WorkerLock(app_env.home)
    assert again.acquire()
    again.release()


@pytest.mark.skipif(os.name == "nt", reason="POSIX signals")
def test_worker_process_lock_and_graceful_stop(app_env):
    """A second `python -m app.worker` exits at once; SIGTERM stops the first cleanly."""
    from app.worker import LOCK_EXIT_CODE

    env = dict(os.environ)
    first = subprocess.Popen([sys.executable, "-m", "app.worker"], cwd=BACKEND, env=env,
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        deadline = time.monotonic() + 30
        while not (app_env.home / "worker.pid").exists() and time.monotonic() < deadline:
            time.sleep(0.1)
        assert (app_env.home / "worker.pid").exists()
        second = subprocess.run([sys.executable, "-m", "app.worker"], cwd=BACKEND, env=env,
                                capture_output=True, timeout=60)
        assert second.returncode == LOCK_EXIT_CODE
        first.send_signal(signal.SIGTERM)
        assert first.wait(timeout=20) == 0
        assert not (app_env.home / "worker.pid").exists()
    finally:
        if first.poll() is None:
            first.kill()


SLOW_SERVER = """#!{python}
import os, sys, time
open({pidfile!r}, "w").write(str(os.getpid()))
sys.path.insert(0, {tests!r})
import fake_llama_server as f
real = f.Handler.do_POST
def slow(self):
    time.sleep(60)
    real(self)
f.Handler.do_POST = slow
f.main()
"""

WORKER_SCRIPT = """
import sys
sys.path.insert(0, {backend!r})
from app import worker
from app.config import get_settings
from app.model_manager import ModelManager
from app.pipeline.translate.llamacpp import LlamaServerEngine

class Manager(ModelManager):
    def create(self, model_id):
        if model_id == "dev-standin-llm":
            return LlamaServerEngine(model_id, "Fake LLM", 0.1, server_bin={server!r}, model_path={model!r})
        return super().create(model_id)

worker.install_signal_handlers()
settings = get_settings()
worker.JobRunner(settings, Manager(settings)).run(worker.next_job_id())
"""


@pytest.mark.skipif(os.name == "nt", reason="POSIX shebang launcher and signals")
def test_sigterm_stops_llama_server_and_keeps_the_job_resumable(client, app_env, pdfs, tmp_path):
    pidfile = tmp_path / "server.pid"
    server = tmp_path / "llama-server"
    server.write_text(SLOW_SERVER.format(python=sys.executable, pidfile=str(pidfile),
                                         tests=str(Path(__file__).parent)))
    server.chmod(server.stat().st_mode | stat.S_IEXEC)
    model = tmp_path / "model.gguf"
    model.write_bytes(b"GGUF")
    script = tmp_path / "run_worker.py"
    script.write_text(WORKER_SCRIPT.format(backend=str(BACKEND), server=str(server), model=str(model)))

    job = upload(client, pdfs["clean"])
    proc = subprocess.Popen([sys.executable, str(script)], cwd=BACKEND, env=dict(os.environ),
                            stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    server_pid = None
    try:
        deadline = time.monotonic() + 60
        while time.monotonic() < deadline and job_row(job["id"]).model_used is None:
            assert proc.poll() is None, proc.stderr.read().decode()
            time.sleep(0.1)
        assert pidfile.exists()
        server_pid = int(pidfile.read_text())
        assert psutil.pid_exists(server_pid)
        time.sleep(1.0)  # now inside a (slow) translation request
        proc.send_signal(signal.SIGTERM)
        assert proc.wait(timeout=30) == 0, proc.stderr.read().decode()
        time.sleep(0.2)
        alive = psutil.pid_exists(server_pid) and psutil.Process(server_pid).status() != psutil.STATUS_ZOMBIE
        assert not alive, "llama-server was left running"
        row = job_row(job["id"])
        assert row.status == "running" and row.stage == "final"  # resumes on the next start
    finally:
        if proc.poll() is None:
            proc.kill()
        if server_pid and psutil.pid_exists(server_pid):
            psutil.Process(server_pid).kill()
